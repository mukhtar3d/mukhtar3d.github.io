"""Import the university database, whatever shape it arrives in.

    python manage.py import_universities --path ../data/universities.csv
    python manage.py import_universities --path provided.sqlite --table unis
    python manage.py import_universities --path unis.json --dry-run

Column names are matched loosely (case, spacing, punctuation and common synonyms
are all ignored), so "Tuition Fee (EUR/year)", "tuition_eur" and "annual tuition"
all land in the same field. Anything unrecognised is preserved in `extra` instead
of being dropped, and --dry-run prints the mapping so you can check it before
committing.
"""

from __future__ import annotations

import csv
import json
import re
import sqlite3
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from api.models import University

# Every alias that should land in a given model field.
FIELD_ALIASES: dict[str, list[str]] = {
    "name": ["name", "university", "universityname", "institution", "school", "title", "uni"],
    "country": ["country", "nation", "countryname", "state"],
    "city": ["city", "town", "location", "campuscity"],
    "website": ["website", "url", "homepage", "web", "link", "site"],
    "description": ["description", "about", "summary", "overview", "notes", "info"],
    "ranking_world": ["rankingworld", "worldrank", "globalrank", "qsrank", "therank", "rank", "ranking", "worldranking"],
    "ranking_national": ["rankingnational", "nationalrank", "countryrank", "localrank"],
    "acceptance_rate": ["acceptancerate", "admissionrate", "admitrate", "acceptance", "selectivity"],
    "tuition_min_eur": ["tuitionmin", "mintuition", "tuitionfrom", "tuitionlow", "minfee"],
    "tuition_max_eur": [
        "tuitionmax", "maxtuition", "tuition", "tuitionfee", "tuitionfees", "fee", "fees",
        "annualtuition", "tuitionperyear", "tuitioneur", "tuitionusd", "cost", "tuitioncost",
    ],
    "living_cost_eur": ["livingcost", "livingcosts", "livingexpenses", "costofliving", "living", "monthlyliving"],
    "housing_cost_eur": ["housingcost", "accommodation", "accommodationcost", "dormcost", "rent", "housing"],
    "application_fee_eur": ["applicationfee", "appfee", "applicationcost"],
    "scholarship_notes": ["scholarship", "scholarships", "financialaid", "aid", "grants", "funding"],
    "scholarship_max_eur": ["scholarshipmax", "maxscholarship", "scholarshipamount"],
    "min_ielts": ["minielts", "ielts", "ieltsrequirement", "ieltsmin", "ieltsscore", "englishrequirement"],
    "avg_ielts": ["avgielts", "averageielts", "meanielts"],
    "min_sat": ["minsat", "sat", "satrequirement", "satmin", "satscore"],
    "avg_sat": ["avgsat", "averagesat", "meansat", "satavg"],
    "sat_required": ["satrequired", "requiressat", "needsat"],
    "min_gpa": ["mingpa", "gpa", "gparequirement", "gpamin", "grade", "minimumgpa"],
    "avg_gpa": ["avggpa", "averagegpa", "meangpa"],
    "language_of_instruction": ["language", "languageofinstruction", "instructionlanguage", "teachinglanguage", "medium"],
    "majors": ["majors", "programs", "programmes", "courses", "fields", "subjects", "faculties", "degrees", "specialisations", "specializations"],
    "application_deadline": ["deadline", "applicationdeadline", "applydeadline", "closingdate", "deadlines"],
    "latitude": ["latitude", "lat"],
    "longitude": ["longitude", "lng", "lon", "long"],
    "campus_map_url": ["campusmap", "campusmapurl", "mapurl", "map"],
    "has_housing": ["hashousing", "dormitory", "dorms", "housingavailable", "accommodationavailable"],
    "image_front": ["imagefront", "image", "photo", "picture", "imageurl", "photourl", "exterior", "frontimage", "thumbnail"],
    "image_inside": ["imageinside", "interior", "insideimage", "interiorimage", "image2"],
    "image_housing": ["imagehousing", "housingimage", "dormimage", "residenceimage", "image3"],
}

NUMERIC_FIELDS = {
    "ranking_world", "ranking_national", "tuition_min_eur", "tuition_max_eur",
    "living_cost_eur", "housing_cost_eur", "application_fee_eur", "scholarship_max_eur",
    "min_sat", "avg_sat",
}
FLOAT_FIELDS = {"acceptance_rate", "min_ielts", "avg_ielts", "min_gpa", "avg_gpa", "latitude", "longitude"}
BOOL_FIELDS = {"sat_required", "has_housing"}

TRUTHY = {"1", "true", "yes", "y", "required", "mandatory", "available", "t"}
FALSY = {"0", "false", "no", "n", "optional", "not required", "none", "f", ""}


def norm_key(key: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(key).lower())


def build_header_map(headers: list[str]) -> tuple[dict[str, str], list[str]]:
    """Map source headers onto model fields. Returns (mapping, unmapped)."""
    lookup: dict[str, str] = {}
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            lookup.setdefault(norm_key(alias), field)

    mapping: dict[str, str] = {}
    unmapped: list[str] = []
    taken: set[str] = set()

    for header in headers:
        key = norm_key(header)
        field = lookup.get(key)
        if field is None:
            # Fall back to substring matching: "tuition_fee_eur_per_year" still
            # finds "tuition".
            for alias_key, candidate in lookup.items():
                if len(alias_key) >= 4 and alias_key in key:
                    field = candidate
                    break
        if field and field not in taken:
            mapping[header] = field
            taken.add(field)
        else:
            unmapped.append(header)
    return mapping, unmapped


def parse_number(value, integer: bool = True):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(round(value)) if integer else float(value)
    text = str(value).strip()
    if not text or text.lower() in {"n/a", "na", "none", "null", "-", "unknown", "tbd"}:
        return None
    text = text.replace("\u202f", "").replace("\xa0", "")
    # Space as a thousands separator ("7 200", "9 000 EUR") is standard in
    # French, Polish, Czech and Scandinavian data. Without this the value gets
    # truncated at the space and €9,000 silently becomes €9.
    text = re.sub(r"(?<=\d)[ \t](?=\d{3}(?:\D|$))", "", text)
    # "1.200,50" (European) vs "1,200.50" (Anglo)
    if re.search(r"\d\.\d{3}(\D|$)", text) and "," in text:
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", "")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    number = float(match.group())
    if "%" in str(value):
        number = number / 100
    return int(round(number)) if integer else number


def parse_bool(value) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in TRUTHY


def parse_list(value) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(v).strip() for v in value if str(v).strip()]
    text = str(value).strip()
    if not text:
        return []
    if text.startswith("["):
        try:
            parsed = json.loads(text)
            if isinstance(parsed, list):
                return [str(v).strip() for v in parsed if str(v).strip()]
        except json.JSONDecodeError:
            pass
    return [part.strip() for part in re.split(r"[;,|/]|\s{2,}", text) if part.strip()]


def coerce(field: str, value):
    if field in NUMERIC_FIELDS:
        return parse_number(value, integer=True)
    if field in FLOAT_FIELDS:
        number = parse_number(value, integer=False)
        if field == "acceptance_rate" and number is not None and number > 1:
            number = number / 100  # "18" meaning 18%
        if field == "min_gpa" and number is not None and number > 4.3:
            number = round(number / (5 if number <= 5 else (10 if number <= 10 else 100)) * 4, 2)
        return number
    if field in BOOL_FIELDS:
        return parse_bool(value)
    if field == "majors":
        return parse_list(value)
    return str(value).strip() if value is not None else ""


# ---------------------------------------------------------------------------
# Readers
# ---------------------------------------------------------------------------
def read_csv(path: Path) -> list[dict]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        sample = handle.read(8192)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return list(csv.DictReader(handle, dialect=dialect))


def read_json(path: Path) -> list[dict]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, dict):
        for key in ("universities", "results", "data", "rows", "items"):
            if isinstance(data.get(key), list):
                return data[key]
        return [data]
    return data


def read_xlsx(path: Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise CommandError("Reading .xlsx needs openpyxl — `pip install openpyxl`.") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(h) if h is not None else f"col{i}" for i, h in enumerate(next(rows))]
    return [dict(zip(headers, row)) for row in rows if any(cell is not None for cell in row)]


def read_sqlite(path: Path, table: str | None) -> list[dict]:
    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    if not table:
        tables = [
            r[0] for r in cursor.execute(
                "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'"
            )
        ]
        if not tables:
            raise CommandError("No tables found in that SQLite file.")
        table = next(
            (t for t in tables if any(k in t.lower() for k in ("uni", "school", "institution"))),
            tables[0],
        )
    rows = [dict(r) for r in cursor.execute(f'SELECT * FROM "{table}"')]
    conn.close()
    return rows


class Command(BaseCommand):
    help = "Import universities from CSV, JSON, XLSX or SQLite with fuzzy column matching."

    def add_arguments(self, parser):
        parser.add_argument("--path", required=True, help="Path to the database file.")
        parser.add_argument("--table", help="Table name (SQLite only). Auto-detected if omitted.")
        parser.add_argument("--truncate", action="store_true", help="Delete existing rows first.")
        parser.add_argument("--dry-run", action="store_true", help="Show the mapping without writing.")
        parser.add_argument(
            "--currency-rate", type=float, default=1.0,
            help="Multiply every money column by this (e.g. 0.92 to convert USD to EUR).",
        )

    def handle(self, *args, **options):
        path = Path(options["path"]).expanduser().resolve()
        if not path.exists():
            raise CommandError(f"No file at {path}")

        suffix = path.suffix.lower()
        if suffix == ".csv" or suffix == ".tsv":
            rows = read_csv(path)
        elif suffix == ".json":
            rows = read_json(path)
        elif suffix in {".xlsx", ".xlsm"}:
            rows = read_xlsx(path)
        elif suffix in {".sqlite", ".sqlite3", ".db"}:
            rows = read_sqlite(path, options.get("table"))
        else:
            raise CommandError(f"Unsupported file type: {suffix}")

        if not rows:
            raise CommandError("That file has no rows.")

        headers = list(rows[0].keys())
        mapping, unmapped = build_header_map(headers)

        self.stdout.write(self.style.MIGRATE_HEADING(f"\n{len(rows)} rows · {len(headers)} columns"))
        for source, field in mapping.items():
            self.stdout.write(f"  {source:<38} → {field}")
        if unmapped:
            self.stdout.write(self.style.WARNING(f"  kept in `extra`: {', '.join(map(str, unmapped))}"))
        if "name" not in mapping.values():
            raise CommandError(
                "No column maps to `name`. Rename the university-name column to "
                "'name' or add it to FIELD_ALIASES."
            )

        if options["dry_run"]:
            self.stdout.write(self.style.SUCCESS("\nDry run — nothing written."))
            self._preview(rows[0], mapping, options["currency_rate"])
            return

        rate = options["currency_rate"]
        created = updated = skipped = 0

        with transaction.atomic():
            if options["truncate"]:
                deleted, _ = University.objects.all().delete()
                self.stdout.write(self.style.WARNING(f"Deleted {deleted} existing rows."))

            for raw in rows:
                values, extra = self._row_to_fields(raw, mapping, rate)
                name = values.get("name", "").strip()
                if not name:
                    skipped += 1
                    continue
                values["extra"] = extra
                values["data_source"] = University.DataSource.PROVIDED

                existing = University.objects.filter(
                    name__iexact=name, city__iexact=values.get("city", "")
                ).first()
                if existing:
                    for field, value in values.items():
                        if value not in (None, "", [], {}):
                            setattr(existing, field, value)
                    existing.save()
                    updated += 1
                else:
                    University.objects.create(**values)
                    created += 1

        self.stdout.write(
            self.style.SUCCESS(
                f"\nDone — {created} created, {updated} updated, {skipped} skipped. "
                f"{University.objects.count()} universities in the database."
            )
        )

    def _row_to_fields(self, raw: dict, mapping: dict, rate: float) -> tuple[dict, dict]:
        values: dict = {}
        extra: dict = {}
        money = {
            "tuition_min_eur", "tuition_max_eur", "living_cost_eur", "housing_cost_eur",
            "application_fee_eur", "scholarship_max_eur",
        }
        for source, value in raw.items():
            field = mapping.get(source)
            if field is None:
                if value not in (None, ""):
                    extra[str(source)] = value if isinstance(value, (int, float, bool)) else str(value)
                continue
            parsed = coerce(field, value)
            if field in money and parsed is not None and rate != 1.0:
                parsed = int(round(parsed * rate))
            if parsed not in (None, ""):
                values[field] = parsed
        return values, extra

    def _preview(self, raw: dict, mapping: dict, rate: float) -> None:
        values, extra = self._row_to_fields(raw, mapping, rate)
        self.stdout.write(self.style.MIGRATE_HEADING("First row as it would be stored:"))
        self.stdout.write(json.dumps(values, indent=2, default=str))
        if extra:
            self.stdout.write(self.style.WARNING("extra: " + json.dumps(extra, default=str)[:400]))
